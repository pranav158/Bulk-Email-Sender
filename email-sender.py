import smtplib
from email.mime.multipart import MIMEMultipart
from email.mime.text import MIMEText
from email.mime.application import MIMEApplication
import openpyxl

def send_bulk_email(smtp_server, smtp_port, smtp_username, smtp_password, sender_email, subject, template_path, excel_path):
    # Read Excel data
    workbook = openpyxl.load_workbook(excel_path)
    sheet = workbook.active

    # Connect to SMTP server with SSL
    with smtplib.SMTP_SSL(smtp_server, smtp_port) as server:
        server.login(smtp_username, smtp_password)

        # Iterate through rows in Excel
        for row in sheet.iter_rows(min_row=2, values_only=True):
            recipient_email = row[0]  # Assuming email is in the first column
            name = row[1]  # Assuming name is in the second column
            user_id = str(row[2])  # Assuming user ID is in the third column and converting to string

            # Read HTML template
            with open(template_path, 'r') as template_file:
                html_template = template_file.read()

            # Replace placeholders with actual data
            html_content = html_template.replace('{name}', name).replace('{id}', user_id)

            # Create email message
            message = MIMEMultipart()
            message['From'] = sender_email
            message['To'] = recipient_email
            message['Subject'] = subject

            # Attach HTML content
            message.attach(MIMEText(html_content, 'html'))

            # Attach PDF file (replace 'pdf_path' with the actual path to your PDF file)
            pdf_path = f'files/{user_id}.pdf'
            with open(pdf_path, 'rb') as pdf_file:
                pdf_attachment = MIMEApplication(pdf_file.read(), 'pdf')
                pdf_attachment.add_header('Content-Disposition', f'attachment; filename={user_id}.pdf')
                message.attach(pdf_attachment)

            # Send email
            server.sendmail(sender_email, recipient_email, message.as_string())

if __name__ == "__main__":
    smtp_server = 'smtp.hostinger.com'
    smtp_port = 465  # Use the appropriate port for SSL (e.g., 465)
    smtp_username = 'admin@edunexa.tech'
    smtp_password = 'Saloni@9283'
    sender_email = 'no-reply@edunexa.tech'
    subject = 'Pog!!!'
    template_path = 'template.html'
    excel_path = 'data.xlsx'

    send_bulk_email(smtp_server, smtp_port, smtp_username, smtp_password, sender_email, subject, template_path, excel_path)