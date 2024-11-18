import smtplib
from email.mime.multipart import MIMEMultipart
from email.mime.text import MIMEText
from email.mime.application import MIMEApplication
import openpyxl
import dkim

def send_bulk_email(smtp_server, smtp_port, smtp_username, smtp_password, sender_email, subject, template_path, excel_path, domain, selector, private_key_path):
    # Read Excel data
    workbook = openpyxl.load_workbook(excel_path)
    sheet = workbook.active

    # Connect to SMTP server with TLS
    with smtplib.SMTP(smtp_server, smtp_port) as server:
        server.starttls()
        server.login(smtp_username, smtp_password)

        # Iterate through rows in Excel
        for row in sheet.iter_rows(min_row=2, values_only=True):
            recipient_email = row[0]  # Assuming email is in the first column
            name = row[1]  # Assuming name is in the second column
            user_id = str(row[2])  # Assuming user ID is in the third column and converting to string

            # Read HTML template
            with open(template_path, 'r') as template_file:
                html_template = template_file.read()

            # Replace placeholders with actual data
            html_content = html_template.replace('{name}', name).replace('{id}', user_id)

            # Create email message
            message = MIMEMultipart()
            message['From'] = sender_email
            message['To'] = recipient_email
            message['Subject'] = subject

            # Attach HTML content
            message.attach(MIMEText(html_content, 'html'))

            # Attach PDF file (replace 'pdf_path' with the actual path to your PDF file)
            pdf_path = f'files/{user_id}.pdf'
            with open(pdf_path, 'rb') as pdf_file:
                pdf_attachment = MIMEApplication(pdf_file.read(), 'pdf')
                pdf_attachment.add_header('Content-Disposition', f'attachment; filename={user_id}.pdf')
                message.attach(pdf_attachment)

            # Sign the email with DKIM
            signed_message = sign_email(message, domain, selector, private_key_path)
            
            # Send email
            server.sendmail(sender_email, recipient_email, signed_message.as_string())

def sign_email(message, domain, selector, private_key_path):
    # Load your private key
    with open(private_key_path, 'rb') as key_file:
        private_key = key_file.read()


    # Convert domain and selector to bytes
    domain_bytes = domain.encode('utf-8')
    selector_bytes = selector.encode('utf-8')


    # Sign the message
    signature = dkim.sign(
        message.as_bytes(),  # Use as_bytes instead of as_string
        domain=domain_bytes,
        selector=selector_bytes,
        privkey=private_key,
        identity=domain_bytes,
        include_headers=["from", "reply-to", "subject", "to", "cc", "mime-version", "content-type"]
    )

    # Add the DKIM signature to the email header
    message['DKIM-Signature'] = signature.decode('utf-8').replace("\r\n", " ")

    return message

if __name__ == "__main__":
    smtp_server = "mail.privateemail.com"
    smtp_port = 587
    smtp_username = "admin@thecyberminds.com"
    smtp_password = "Pk.9527799"
    sender_email = "admin@thecyberminds.com"
    sender_name = "TheCyberMinds"
    subject = "Test"
    template_path = 'template.html'
    excel_path = 'data.xlsx'

    # DKIM configuration
    domain = 'thecyberminds.com'  # Replace with your domain
    selector = 'default'  # Replace with your DKIM selector
    private_key_path = 'thecyberminds.pem'  # Replace with the path to your private key file

    send_bulk_email(smtp_server, smtp_port, smtp_username, smtp_password, sender_email, subject, template_path, excel_path, domain, selector, private_key_path)
